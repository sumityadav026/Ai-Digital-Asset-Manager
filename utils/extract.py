import re
from PyPDF2 import PdfReader
from docx import Document


def _clean_text(text: str) -> str:
    """Normalise whitespace while preserving newlines that separate sections."""
    # Collapse multiple spaces/tabs into one (but not newlines)
    text = re.sub(r"[ \t]+", " ", text)
    # Collapse 3+ consecutive newlines into 2
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(file_path: str) -> str:
    """
    Extract plain text from PDF or DOCX.

    For PDFs each page is extracted with PyPDF2 and separated by a newline so
    that heading/section breaks at page boundaries are preserved.  The caller
    is responsible for chunking.
    """
    text = ""

    try:
        if file_path.lower().endswith(".pdf"):
            reader = PdfReader(file_path)
            page_texts = []
            for page in reader.pages:
                content = page.extract_text()
                if content:
                    page_texts.append(content)
            # Join pages with a separator so cross-page boundaries are clear
            text = "\n\n".join(page_texts)

        elif file_path.lower().endswith(".docx"):
            doc = Document(file_path)
            parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    parts.append(para.text)
            # Also pull text from tables (important for CVs/invoices)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        parts.append(row_text)
            text = "\n".join(parts)

    except Exception as e:
        print(f"Error extracting text from {file_path}: {e}")

    return _clean_text(text)